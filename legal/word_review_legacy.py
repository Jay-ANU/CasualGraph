"""Paragraph-level Word (.docx) review helpers carried over from the retired desktop companion.

Plain functions only: no FastAPI, database, retrieval or model-client dependencies, so
drafting code can reuse them. The writing goal/template tables were reduced to a single
neutral "general" review instruction.
"""

from __future__ import annotations

import io
import json
import re
from pathlib import Path
from typing import Any, Dict, List, Mapping, Optional, Tuple
from urllib.parse import unquote

from docx import Document as DocxDocument

WORD_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# Former DESKTOP_WORD_EDIT_* defaults.
DEFAULT_MAX_PARAGRAPHS = 28
DEFAULT_MAX_CHARS = 14_000
DEFAULT_MAX_SUGGESTIONS = 8

GENERAL_REVIEW_LABEL = "General review"
GENERAL_REVIEW_INSTRUCTION = "Review the draft without forcing a specific document structure."

WORD_EDIT_CATEGORY_LABELS = {
    "clarity": "Clarity",
    "logic": "Logic",
    "evidence": "Evidence",
    "structure": "Structure",
    "tone": "Tone",
    "esg_concept": "ESG concept",
}

WORD_EDIT_EVIDENCE_GAP_LABELS = {
    "metric": "metric",
    "source": "source",
    "comparison": "comparison",
    "causal_link": "causal link",
    "citation": "citation",
    "concept_definition": "concept definition",
}


def safe_word_filename(filename: Optional[str]) -> str:
    name = Path(unquote(str(filename or "document.docx"))).name
    name = re.sub(r"[^A-Za-z0-9._ ()-]+", "_", name).strip(" ._") or "document.docx"
    if not name.lower().endswith(".docx"):
        name = f"{Path(name).stem or 'document'}.docx"
    return name


def unique_output_name(file_name: str) -> str:
    stem = Path(file_name).stem or "document"
    return safe_word_filename(f"{stem}.edited.docx")


def parse_docx_paragraphs(file_bytes: bytes) -> Tuple[Any, List[Dict[str, Any]]]:
    """Return the python-docx document and its non-empty paragraphs as ``p_001``-style records.

    Raises ValueError when the document has no extractable paragraph text.
    """
    document = DocxDocument(io.BytesIO(file_bytes))
    paragraphs: List[Dict[str, Any]] = []
    visible_index = 0
    for docx_index, paragraph in enumerate(document.paragraphs):
        text = re.sub(r"\s+", " ", str(paragraph.text or "")).strip()
        if not text:
            continue
        visible_index += 1
        paragraphs.append({
            "id": f"p_{visible_index:03d}",
            "docx_index": docx_index,
            "text": text,
        })
    if not paragraphs:
        raise ValueError("The Word document does not contain extractable paragraph text.")
    return document, paragraphs


def select_word_paragraphs_for_review(
    paragraphs: List[Dict[str, Any]],
    *,
    max_chars: int = DEFAULT_MAX_CHARS,
    max_paragraphs: int = DEFAULT_MAX_PARAGRAPHS,
) -> List[Dict[str, str]]:
    selected: List[Dict[str, str]] = []
    total_chars = 0
    for paragraph in paragraphs:
        text = str(paragraph.get("text") or "").strip()
        if len(text) < 35 and len(selected) >= 3:
            continue
        clipped = text[:1400]
        if selected and total_chars + len(clipped) > max_chars:
            break
        selected.append({"id": str(paragraph.get("id") or ""), "text": clipped})
        total_chars += len(clipped)
        if len(selected) >= max_paragraphs:
            break
    return selected or [{"id": str(paragraphs[0]["id"]), "text": str(paragraphs[0]["text"])[:1400]}]


def build_word_edit_messages(
    instruction: str,
    paragraphs: List[Dict[str, Any]],
    *,
    evidence_sources: Optional[List[Dict[str, Any]]] = None,
    max_suggestions: int = DEFAULT_MAX_SUGGESTIONS,
    max_chars: int = DEFAULT_MAX_CHARS,
    max_paragraphs: int = DEFAULT_MAX_PARAGRAPHS,
) -> List[Dict[str, str]]:
    """Chat messages asking a model for paragraph-level replacement suggestions (JSON only)."""
    review_paragraphs = select_word_paragraphs_for_review(
        paragraphs,
        max_chars=max_chars,
        max_paragraphs=max_paragraphs,
    )
    prompt = {
        "writing_template": GENERAL_REVIEW_LABEL,
        "template_instruction": GENERAL_REVIEW_INSTRUCTION,
        "task": str(instruction or "").strip()
        or "Improve this Word document for academic or business analysis while preserving factual meaning.",
        "rules": [
            "Return JSON only.",
            "Suggest paragraph-level replacements only; do not invent facts, data, or citations.",
            "Preserve the user's meaning and named entities.",
            "Prioritize clarity, structure, analytical strength, ESG/business terminology, and evidence-aware phrasing.",
            "Each suggestion must include category: clarity, logic, evidence, structure, tone, or esg_concept.",
            "Each suggestion must state the concrete problem it solves for the writer.",
            "Each suggestion must include severity: low, medium, or high.",
            "Use evidence_refs only when an evidence source directly supports the change.",
            "If a paragraph needs evidence but no source supports it, set evidence_needed true and explain the gap.",
            "Use evidence_gap_types to flag missing metric, source, comparison, causal_link, citation, or concept_definition.",
            f"Return at most {max_suggestions} suggestions.",
        ],
        "schema": {
            "suggestions": [
                {
                    "paragraph_id": "p_001",
                    "category": "clarity|logic|evidence|structure|tone|esg_concept",
                    "severity": "low|medium|high",
                    "problem": "The concrete weakness this edit fixes.",
                    "replacement": "Full replacement paragraph text.",
                    "reason": "Why this edit helps.",
                    "evidence_refs": ["E1"],
                    "evidence_needed": False,
                    "evidence_gap_types": ["metric", "source", "comparison", "causal_link", "citation", "concept_definition"],
                }
            ]
        },
        "evidence_sources": evidence_sources or [],
        "paragraphs": review_paragraphs,
    }
    return [
        {
            "role": "system",
            "content": (
                "You are a careful Word document editor for business, ESG, finance, and academic writing. "
                "You help users improve draft quality, but you must not fabricate evidence."
            ),
        },
        {"role": "user", "content": json.dumps(prompt, ensure_ascii=False)},
    ]


def normalize_word_edit_suggestions(
    parsed: Dict[str, Any],
    paragraph_lookup: Dict[str, Dict[str, Any]],
    evidence_sources: Optional[List[Dict[str, Any]]] = None,
    *,
    max_suggestions: int = DEFAULT_MAX_SUGGESTIONS,
) -> List[Dict[str, Any]]:
    raw_suggestions = parsed.get("suggestions")
    if not isinstance(raw_suggestions, list):
        return []
    allowed_evidence_ids = {str(item.get("id") or "") for item in (evidence_sources or [])}
    suggestions: List[Dict[str, Any]] = []
    seen_paragraphs: set[str] = set()
    for raw_item in raw_suggestions:
        if not isinstance(raw_item, dict):
            continue
        paragraph_id = str(raw_item.get("paragraph_id") or raw_item.get("paragraph") or "").strip()
        if paragraph_id not in paragraph_lookup or paragraph_id in seen_paragraphs:
            continue
        replacement = re.sub(r"\s+", " ", str(raw_item.get("replacement") or "")).strip()
        if not replacement:
            continue
        original = str(paragraph_lookup[paragraph_id].get("text") or "").strip()
        if replacement == original:
            continue
        reason = re.sub(r"\s+", " ", str(raw_item.get("reason") or "Improves clarity and analytical quality.")).strip()
        problem = re.sub(
            r"\s+",
            " ",
            str(raw_item.get("problem") or raw_item.get("problem_solved") or "Improves draft quality.").strip(),
        )
        category = str(raw_item.get("category") or "clarity").strip().lower()
        if category not in WORD_EDIT_CATEGORY_LABELS:
            category = "clarity"
        severity = str(raw_item.get("severity") or "").strip().lower()
        if severity not in {"low", "medium", "high"}:
            severity = "medium" if category in {"logic", "evidence", "esg_concept"} else "low"
        raw_refs = raw_item.get("evidence_refs") if isinstance(raw_item.get("evidence_refs"), list) else []
        evidence_refs = [str(ref) for ref in raw_refs if str(ref) in allowed_evidence_ids]
        raw_gaps = raw_item.get("evidence_gap_types") if isinstance(raw_item.get("evidence_gap_types"), list) else []
        evidence_gap_types = [
            str(gap).strip().lower()
            for gap in raw_gaps
            if str(gap).strip().lower() in WORD_EDIT_EVIDENCE_GAP_LABELS
        ]
        evidence_needed = bool(raw_item.get("evidence_needed")) or (category == "evidence" and not evidence_refs)
        if evidence_needed and not evidence_gap_types:
            evidence_gap_types = ["source"]
        suggestions.append({
            "id": f"s_{len(suggestions) + 1:03d}",
            "paragraph_id": paragraph_id,
            "operation": "replace",
            "category": category,
            "category_label": WORD_EDIT_CATEGORY_LABELS[category],
            "severity": severity,
            "original": original,
            "replacement": replacement,
            "problem": problem[:260],
            "reason": reason[:500],
            "evidence_refs": evidence_refs,
            "evidence_needed": evidence_needed,
            "evidence_gap_types": evidence_gap_types,
            "sources": [paragraph_id],
        })
        seen_paragraphs.add(paragraph_id)
        if len(suggestions) >= max_suggestions:
            break
    return suggestions


def fallback_word_edit_suggestions(paragraphs: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    suggestions: List[Dict[str, Any]] = []
    for paragraph in paragraphs:
        text = str(paragraph.get("text") or "").strip()
        normalized = re.sub(r"\s+", " ", text).strip()
        if normalized and normalized != text:
            suggestions.append({
                "id": f"s_{len(suggestions) + 1:03d}",
                "paragraph_id": str(paragraph.get("id") or ""),
                "operation": "replace",
                "category": "clarity",
                "category_label": WORD_EDIT_CATEGORY_LABELS["clarity"],
                "severity": "low",
                "original": text,
                "replacement": normalized,
                "problem": "The paragraph contains inconsistent spacing that makes the draft less polished.",
                "reason": "Normalizes spacing without changing the meaning.",
                "evidence_refs": [],
                "evidence_needed": False,
                "evidence_gap_types": [],
                "sources": [str(paragraph.get("id") or "")],
            })
        if len(suggestions) >= 3:
            break
    return suggestions


def replace_paragraph_text(paragraph: Any, replacement: str) -> None:
    text = str(replacement or "")
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
        return
    paragraph.add_run(text)


def apply_paragraph_replacements(document: Any, replacements: Mapping[str, str]) -> int:
    """Replace paragraphs by ``p_001``-style id (same numbering as parse_docx_paragraphs).

    Returns how many paragraphs were changed.
    """
    applied_count = 0
    visible_index = 0
    for paragraph in document.paragraphs:
        if not str(paragraph.text or "").strip():
            continue
        visible_index += 1
        paragraph_id = f"p_{visible_index:03d}"
        replacement = replacements.get(paragraph_id)
        if replacement:
            replace_paragraph_text(paragraph, replacement)
            applied_count += 1
    return applied_count

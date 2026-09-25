"""Runtime helpers for ingesting documents: dedup, parse, clean, chunk, index, register."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional
import io
import json
import re
import shutil

from docx import Document as DocxDocument
from pypdf import PdfReader

import document_registry
from configs.settings import (
    ACTIVE_VECTOR_STORE_FILE,
    CHUNK_DIR,
    PINECONE_NAMESPACE,
    PROCESSED_DIR,
    VECTOR_DIR,
    VECTOR_STORE_PROVIDER,
    ensure_directories,
)
from document_processing.chunker import chunk_text
from document_processing.text_cleaner import clean_text
from graph.neo4j_store import get_neo4j_store
from rag.bm25_index import build_bm25_index
from rag.pinecone_store import delete_vectors_by_document_id, pinecone_available
import rag.vector_store as vector_store_module
from rag.vector_store import build_vector_store


def ingest_uploaded_document(
    title: str,
    domain: str = "general",
    source: str = "",
    content: str = "",
    document_group: str = "user_upload",
    source_type: str = "",
    owner_user_id: str = "",
    visibility_scope: str = "global",
    filename: str | None = None,
    file_bytes: bytes | None = None,
    file_path: str | None = None,
    raw_hash: str | None = None,
    progress_callback: Optional[Callable[[str, str, int], None]] = None,
) -> Dict:
    """Parse uploaded content, append it to the corpus, and return frontend-friendly data."""
    ensure_directories()
    _report_progress(progress_callback, "reading", "Reading uploaded content", 5)
    owner_value = str(owner_user_id or "").strip()
    scope_value = str(visibility_scope or "global").strip().lower()
    if scope_value not in {"global", "private"}:
        scope_value = "global"

    raw_hash = raw_hash or (
        document_registry.compute_file_hash(file_path)
        if file_path
        else document_registry.compute_raw_hash(file_bytes, content)
    )
    if raw_hash:
        existing = document_registry.lookup(
            raw_hash=raw_hash,
            text_hash=None,
            document_group=document_group,
            owner_user_id=owner_value,
        )
        if existing:
            return _build_duplicate_response(existing, progress_callback=progress_callback)

    text, detected_source = _resolve_text_input(
        content=content,
        filename=filename,
        file_bytes=file_bytes,
        file_path=file_path,
    )
    _report_progress(progress_callback, "cleaning", "Cleaning extracted text", 15)
    cleaned = clean_text(text)
    text_hash = document_registry.compute_text_hash(cleaned)
    existing = document_registry.lookup(
        raw_hash=None,
        text_hash=text_hash,
        document_group=document_group,
        owner_user_id=owner_value,
    )
    if existing:
        return _build_duplicate_response(existing, progress_callback=progress_callback)

    slug = _slugify(title or filename or "document")
    timestamp = datetime.utcnow().strftime("%Y%m%d%H%M%S")
    name = f"{slug}_{timestamp}"
    source_value = source or detected_source
    source_type_value = source_type or ("uploaded_file" if filename else "manual_input")

    _report_progress(progress_callback, "chunking", "Splitting document into chunks", 25)
    chunks = chunk_text(cleaned)
    if not chunks:
        raise ValueError("No usable text chunks were created from the uploaded content.")
    chunks = [
        {
            **chunk,
            "document_id": name,
            "document_title": title or filename or "Untitled document",
            "document_group": document_group,
            "owner_user_id": owner_value,
            "visibility_scope": scope_value,
            "source_type": source_type_value,
            "domain": domain,
            "source": source_value,
        }
        for chunk in chunks
    ]

    processed_text_path = PROCESSED_DIR / f"{name}.txt"
    chunks_path = CHUNK_DIR / f"{name}_chunks.jsonl"
    vector_store_path = VECTOR_DIR / name

    processed_text_path.write_text(cleaned, encoding="utf-8")
    with chunks_path.open("w", encoding="utf-8") as handle:
        for chunk in chunks:
            handle.write(json.dumps(chunk, ensure_ascii=False) + "\n")

    _report_progress(progress_callback, "embedding", "Building vector index", 40)
    build_vector_store(chunks, str(vector_store_path))
    build_bm25_index(chunks, str(vector_store_path))

    # Ingestion no longer extracts entities or builds a graph, so nothing is synced to Neo4j.
    neo4j_sync = {"enabled": False, "synced": False, "reason": "not_synced_on_ingest"}
    _report_progress(progress_callback, "completed", "Document processing complete", 100)

    document_registry.register({
        "document_id": name,
        "title": title or filename or "Untitled document",
        "domain": domain,
        "source": source_value,
        "source_type": source_type_value,
        "document_group": document_group,
        "owner_user_id": owner_value,
        "visibility_scope": scope_value,
        "raw_hash": raw_hash,
        "text_hash": text_hash,
        "ingested_at": datetime.now(timezone.utc).isoformat(),
        "paths": {
            "processed_text": str(processed_text_path),
            "chunks": str(chunks_path),
            "vector_store": str(vector_store_path),
        },
    })

    return {
        "document": {
            "id": name,
            "title": title or filename or "Untitled document",
            "domain": domain,
            "source": source_value,
            "document_group": document_group,
            "owner_user_id": owner_value,
            "visibility_scope": scope_value,
            "source_type": source_type_value,
            "content_hash": text_hash,
            "neo4j_sync": neo4j_sync,
        },
        "stats": {
            "chunk_count": len(chunks),
        },
        "neo4j": neo4j_sync,
    }


def delete_uploaded_document(upload: Dict[str, Any]) -> Dict[str, Any]:
    """Best-effort cleanup for a locally uploaded document."""
    if str(upload.get("status") or "") == "rejected":
        return {
            "deleted_paths": [],
            "warnings": [],
            "neo4j": {"enabled": False, "deleted": False, "reason": "rejected_upload_has_no_resources"},
        }

    deleted_paths: List[str] = []
    warnings: List[str] = []
    paths = upload.get("paths") or {}
    document_id = str(upload.get("document_id") or "").strip()

    for key in ("processed_text_path", "chunks_path", "extractions_path", "graph_path"):
        value = str(paths.get(key) or "").strip()
        if not value:
            continue
        path = Path(value)
        try:
            if path.exists() and path.is_file():
                path.unlink()
                deleted_paths.append(str(path))
        except Exception as exc:
            warnings.append(f"{key}: {type(exc).__name__}: {exc}")

    vector_path = str(paths.get("vector_store_path") or "").strip()
    if vector_path:
        path = Path(vector_path)
        if VECTOR_STORE_PROVIDER.strip().lower() == "pinecone" and document_id:
            try:
                if pinecone_available():
                    delete_vectors_by_document_id(document_id, namespace=PINECONE_NAMESPACE)
                else:
                    warnings.append("pinecone: SDK unavailable; remote vectors were not deleted")
            except Exception as exc:
                warnings.append(f"pinecone: {type(exc).__name__}: {exc}")
        try:
            if path.exists() and path.is_dir():
                shutil.rmtree(path)
                deleted_paths.append(str(path))
                _repair_active_vector_store_after_delete(path)
        except Exception as exc:
            warnings.append(f"vector_store_path: {type(exc).__name__}: {exc}")

    neo4j_result: Dict[str, Any] = {"enabled": False, "deleted": False, "reason": "not_attempted"}
    if document_id:
        try:
            store = get_neo4j_store()
            if store is None:
                neo4j_result = {"enabled": False, "deleted": False, "reason": "neo4j_unavailable"}
            else:
                neo4j_result = store.delete_document(document_id)
        except Exception as exc:
            neo4j_result = {"enabled": True, "deleted": False, "reason": f"{type(exc).__name__}: {exc}"}

    return {"deleted_paths": deleted_paths, "warnings": warnings, "neo4j": neo4j_result}


def summarize_registered_document(entry: Dict[str, Any], audit: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    """List-view payload. Server paths, graph and relationship data stay out of API responses."""
    stats = (audit or {}).get("stats") or {}
    return _document_payload(entry, chunk_count=int(stats.get("chunks") or 0))


def load_registered_document(entry: Dict[str, Any], audit: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    """Detail payload; the chunk count is read from the stored chunk file."""
    chunks_path = Path(str((entry.get("paths") or {}).get("chunks") or ""))
    chunks = _load_jsonl_rows(chunks_path) if chunks_path.is_file() else []
    return _document_payload(entry, chunk_count=len(chunks))


def _document_payload(entry: Dict[str, Any], *, chunk_count: int) -> Dict[str, Any]:
    # Older registry entries may still carry graph/extraction paths; they are never echoed.
    return {
        "id": str(entry.get("document_id") or "").strip(),
        "title": entry.get("title", ""),
        "domain": str(entry.get("domain") or "general"),
        "source": entry.get("source", ""),
        "document_group": entry.get("document_group", ""),
        "owner_user_id": entry.get("owner_user_id", ""),
        "visibility_scope": entry.get("visibility_scope", "global"),
        "source_type": entry.get("source_type", ""),
        "chunk_count": chunk_count,
        "neo4j_sync": dict(entry.get("neo4j_sync") or {}),
        "ingested_at": entry.get("ingested_at") or "",
    }


def _repair_active_vector_store_after_delete(deleted_path: Path) -> None:
    try:
        if not ACTIVE_VECTOR_STORE_FILE.exists():
            return
        raw = ACTIVE_VECTOR_STORE_FILE.read_text(encoding="utf-8").strip()
        manifest = json.loads(raw) if raw.startswith("{") else {"location": raw}
        active_location = Path(str(manifest.get("location") or "")).resolve()
        if active_location != deleted_path.resolve():
            return

        candidates = [
            path
            for path in VECTOR_DIR.iterdir()
            if path.is_dir() and path.resolve() != deleted_path.resolve() and (path / "metadata.json").exists()
        ]
        candidates.sort(key=lambda item: item.stat().st_mtime, reverse=True)
        if candidates:
            replacement = candidates[0].resolve()
            ACTIVE_VECTOR_STORE_FILE.write_text(
                json.dumps(
                    {
                        "provider": VECTOR_STORE_PROVIDER.strip().lower() or "local",
                        "location": str(replacement),
                        "namespace": PINECONE_NAMESPACE,
                        "document_id": replacement.name,
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
        else:
            ACTIVE_VECTOR_STORE_FILE.unlink(missing_ok=True)
        vector_store_module._LOADED_STORE = None
        vector_store_module._LOADED_STORE_KEY = None
    except Exception as exc:
        print(f"[ingestion] Active vector store repair failed: {type(exc).__name__}: {exc}")


def _build_duplicate_response(
    entry: Dict[str, Any],
    progress_callback: Optional[Callable[[str, str, int], None]] = None,
) -> Dict[str, Any]:
    matched_by = str(entry.get("matched_by") or "text_hash")
    paths = entry.get("paths") or {}
    document_id = str(entry.get("document_id") or "")
    domain = str(entry.get("domain") or "general")

    chunks_path = Path(str(paths.get("chunks") or ""))
    chunks = _load_jsonl_rows(chunks_path) if chunks_path.is_file() else []

    print(f"[ingestion] Duplicate detected (matched_by={matched_by}) document_id={document_id}")
    _report_progress(
        progress_callback,
        "completed",
        "Duplicate detected; reusing existing document",
        100,
    )

    return {
        "duplicate": True,
        "matched_by": matched_by,
        "document": {
            "id": document_id,
            "title": entry.get("title", ""),
            "domain": domain,
            "source": entry.get("source", ""),
            "document_group": entry.get("document_group", ""),
            "owner_user_id": entry.get("owner_user_id", ""),
            "visibility_scope": entry.get("visibility_scope", "global"),
            "source_type": entry.get("source_type", ""),
            "content_hash": entry.get("text_hash") or "",
            "neo4j_sync": {"enabled": False, "synced": False, "reason": "duplicate_skipped"},
        },
        "stats": {
            "chunk_count": len(chunks),
        },
        "neo4j": {"enabled": False, "synced": False, "reason": "duplicate_skipped"},
    }


def _resolve_text_input(
    content: str = "",
    filename: str | None = None,
    file_bytes: bytes | None = None,
    file_path: str | None = None,
) -> tuple[str, str]:
    if content and content.strip():
        return content.strip(), "manual_input"

    path = Path(file_path) if file_path else None
    if not filename or (file_bytes is None and path is None):
        raise ValueError("Either non-empty content or an uploaded file is required.")

    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        if path is not None:
            with path.open("rb") as handle:
                reader = PdfReader(handle)
                parts = []
                for page in reader.pages:
                    page_text = (page.extract_text() or "").strip()
                    if page_text:
                        parts.append(page_text)
        else:
            reader = PdfReader(io.BytesIO(file_bytes or b""))
            parts = []
            for page in reader.pages:
                page_text = (page.extract_text() or "").strip()
                if page_text:
                    parts.append(page_text)
        text = "\n\n".join(parts)
        if not text.strip():
            raise ValueError("No extractable text found in the uploaded PDF.")
        return text.strip(), filename

    if suffix in {".txt", ".md", ".markdown"}:
        if path is not None:
            return path.read_text(encoding="utf-8", errors="ignore").strip(), filename
        return (file_bytes or b"").decode("utf-8", errors="ignore").strip(), filename

    if suffix in {".docx", ".doc"}:
        doc = DocxDocument(str(path)) if path is not None else DocxDocument(io.BytesIO(file_bytes or b""))
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())
        if not text.strip():
            raise ValueError("No extractable text found in the uploaded Word document.")
        return text.strip(), filename

    if suffix == ".rtf":
        if path is not None:
            raw = path.read_text(encoding="utf-8", errors="ignore")
        else:
            raw = (file_bytes or b"").decode("utf-8", errors="ignore")
        text = re.sub(r"\\[a-z]+\d*\s?", "", raw)
        text = re.sub(r"\{[^}]*\}", "", text)
        text = re.sub(r"\s+", " ", text).strip()
        if not text:
            raise ValueError("No extractable text found in the uploaded RTF file.")
        return text, filename

    raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")


def _load_jsonl_rows(path: Path) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for raw_line in handle:
            line = raw_line.strip()
            if not line:
                continue
            parsed = json.loads(line)
            if isinstance(parsed, dict):
                rows.append(parsed)
    return rows


def _slugify(text: str) -> str:
    value = re.sub(r"[^a-zA-Z0-9]+", "_", text.strip().lower())
    return value.strip("_") or "document"


def _report_progress(
    callback: Optional[Callable[[str, str, int], None]],
    stage: str,
    message: str,
    progress: int,
) -> None:
    if callback is None:
        return
    callback(stage, message, max(0, min(progress, 100)))

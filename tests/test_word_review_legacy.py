import io
import json

import pytest
from docx import Document

from legal import word_review_legacy as word


def _docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("1. The Supplier shall deliver the Goods by 1 March.")
    document.add_paragraph("   ")
    document.add_paragraph("2. Payment   is due within 30 days.")
    multi_run = document.add_paragraph("3. Either party ")
    multi_run.add_run("may terminate ")
    multi_run.add_run("on notice.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_round_trip_parse_replace_save():
    document, paragraphs = word.parse_docx_paragraphs(_docx_bytes())

    assert [item["id"] for item in paragraphs] == ["p_001", "p_002", "p_003"]
    assert [item["docx_index"] for item in paragraphs] == [0, 2, 3]
    assert paragraphs[1]["text"] == "2. Payment is due within 30 days."
    assert paragraphs[2]["text"] == "3. Either party may terminate on notice."

    applied = word.apply_paragraph_replacements(
        document,
        {"p_002": "2. Payment is due within 45 days.", "p_003": "3. Either party may terminate on 30 days' notice."},
    )
    assert applied == 2

    buffer = io.BytesIO()
    document.save(buffer)
    _, reparsed = word.parse_docx_paragraphs(buffer.getvalue())

    assert [item["text"] for item in reparsed] == [
        "1. The Supplier shall deliver the Goods by 1 March.",
        "2. Payment is due within 45 days.",
        "3. Either party may terminate on 30 days' notice.",
    ]


def test_replace_paragraph_text_keeps_first_run_and_clears_the_rest():
    document, paragraphs = word.parse_docx_paragraphs(_docx_bytes())
    paragraph = document.paragraphs[paragraphs[2]["docx_index"]]

    word.replace_paragraph_text(paragraph, "Replaced clause.")

    assert paragraph.text == "Replaced clause."
    assert [run.text for run in paragraph.runs] == ["Replaced clause.", "", ""]


def test_empty_document_is_rejected_without_http_types():
    buffer = io.BytesIO()
    Document().save(buffer)

    with pytest.raises(ValueError):
        word.parse_docx_paragraphs(buffer.getvalue())


def test_output_names_are_sanitised():
    assert word.safe_word_filename("../drafts/Master Agreement?.txt") == "Master Agreement_.docx"
    assert word.safe_word_filename(None) == "document.docx"
    assert word.unique_output_name("contract.docx") == "contract.edited.docx"


def test_suggestions_are_normalised_against_known_paragraphs():
    _, paragraphs = word.parse_docx_paragraphs(_docx_bytes())
    lookup = {item["id"]: item for item in paragraphs}
    parsed = {
        "suggestions": [
            {"paragraph_id": "p_002", "replacement": "2. Payment is due within 45 days.", "category": "unknown", "evidence_refs": ["E1", "E9"]},
            {"paragraph_id": "p_002", "replacement": "duplicate paragraph is ignored"},
            {"paragraph_id": "p_404", "replacement": "unknown paragraph is ignored"},
            {"paragraph_id": "p_001", "replacement": paragraphs[0]["text"]},
        ]
    }

    suggestions = word.normalize_word_edit_suggestions(parsed, lookup, evidence_sources=[{"id": "E1"}])

    assert len(suggestions) == 1
    assert suggestions[0]["id"] == "s_001"
    assert suggestions[0]["category"] == "clarity"
    assert suggestions[0]["evidence_refs"] == ["E1"]
    assert suggestions[0]["original"] == "2. Payment is due within 30 days."


def test_prompt_uses_the_single_general_instruction():
    _, paragraphs = word.parse_docx_paragraphs(_docx_bytes())

    messages = word.build_word_edit_messages("Tighten the payment clause.", paragraphs)

    assert [message["role"] for message in messages] == ["system", "user"]
    prompt = json.loads(messages[1]["content"])
    assert prompt["template_instruction"] == word.GENERAL_REVIEW_INSTRUCTION
    assert "writing_goal" not in prompt
    assert prompt["task"] == "Tighten the payment clause."
    assert [item["id"] for item in prompt["paragraphs"]] == ["p_001", "p_002", "p_003"]

"""Real files through intake, redaction, model context and user-controlled exports."""
from copy import deepcopy
import io
import json
import sqlite3
import zipfile

from cryptography.fernet import Fernet
from docx import Document
from docx.oxml import OxmlElement
from lxml import etree
from PIL import Image
from pypdf import PdfReader
import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient

from api.routers import contract_review as routes
from legal import contract_documents as documents, document_context as context, document_ocr, review_chat, review_store as store
from legal.docx_redlines import NS


def png(width=360, height=180):
    stream = io.BytesIO()
    Image.new('RGB', (width, height), '#375369').save(stream, format='PNG')
    return stream.getvalue()


def word():
    doc = Document()
    doc.add_paragraph('甲方应在验收后90日内支付价款。')
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = '联系人：'
    table.cell(0, 1).text = '张三'
    table.cell(1, 0).text = '金额'
    table.cell(1, 1).text = '100万元'
    table.cell(2, 0).merge(table.cell(2, 1)).text = '本项跨两列。'
    picture = doc.add_picture(io.BytesIO(png()))
    picture._inline.docPr.set('descr', 'PRIVATE ALT 张三 private@example.test')
    doc.sections[0].header.paragraphs[0].text = '通知应以书面方式送达。'
    result = io.BytesIO()
    doc.save(result)
    return result.getvalue()


def pdf(*, scan_only=False):
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    stream = io.BytesIO()
    drawing = canvas.Canvas(stream, pagesize=(420, 590))
    if not scan_only:
        drawing.drawString(32, 550, 'Buyer shall pay in 90 days after acceptance.')
    drawing.drawImage(ImageReader(io.BytesIO(png())), 40, 180, 240, 120)
    drawing.showPage()
    drawing.save()
    return stream.getvalue()


@pytest.fixture
def api(tmp_path, monkeypatch):
    path = tmp_path / 'structured.db'
    cipher = Fernet(Fernet.generate_key())

    def connect():
        conn = sqlite3.connect(path)
        conn.row_factory = sqlite3.Row
        conn.execute('CREATE TABLE IF NOT EXISTS org_members (org_id TEXT,user_id TEXT,role TEXT)')
        conn.execute('''CREATE TABLE IF NOT EXISTS audit_events(org_id TEXT,matter_id TEXT,actor_user_id TEXT,
                        action TEXT,target_type TEXT,target_id TEXT,details_json TEXT,created_at TEXT)''')
        conn.commit()
        return conn

    monkeypatch.setattr(store, 'connect', connect)
    monkeypatch.setattr(store, 'encode', lambda value: cipher.encrypt(json.dumps(value, ensure_ascii=False).encode()))
    monkeypatch.setattr(store, 'decode', lambda value: json.loads(cipher.decrypt(value)))
    with connect() as conn:
        conn.execute("INSERT INTO org_members VALUES ('o1','u1','owner')")
    actor = {'id': 'u1', 'matter_role': 'member'}

    def member(mid, user, min_role='viewer', require_active=False):
        if (user['id'] != 'u1' or mid != 'm1'
                or (min_role == 'member' and user['matter_role'] == 'viewer')):
            raise HTTPException(403, 'no access')
        return {'matter': {'id': 'm1', 'org_id': 'o1', 'status': 'active'}, 'role': user['matter_role']}

    monkeypatch.setattr(routes.matters, 'require_matter_member', member)
    monkeypatch.setattr(routes.ydata, 'select_model', lambda mid: {'id': mid, 'provider': 'anthropic', 'family': 'Claude'})
    monkeypatch.setattr(routes.ydata, 'chat_json', lambda *a, **kw: pytest.fail('local operation called a model'))
    monkeypatch.setattr(routes.ydata, 'chat', lambda *a, **kw: pytest.fail('local operation called a model'))
    monkeypatch.setattr(routes.engine, 'submit', lambda rid: None)
    monkeypatch.delenv('LEGAL_DOCUMENT_PARSER_URL', raising=False)
    monkeypatch.delenv('LEGAL_DOCUMENT_PARSER_TOKEN', raising=False)
    app = FastAPI()
    app.include_router(routes.router)
    app.include_router(routes.public_router)
    app.dependency_overrides[routes.get_current_user] = lambda: actor
    app.dependency_overrides[routes.require_legal_max] = lambda: actor
    return TestClient(app), actor, path


def upload(client, data, ext):
    response = client.post('/legal/contracts', data={'matter_id': 'm1', 'original_upload_confirmed': 'true',
                                                     'upload_notice_version': 'original-upload-v1'},
                           files={'file': ('synthetic.' + ext, data)})
    assert response.status_code == 201, response.text
    return response.json()


def start(client, contract, tier='ultra_fast'):
    cid = contract['id']
    assert client.post(f'/legal/contracts/{cid}/redaction', json={'revision': 1, 'confirmed': True}).status_code == 200
    block = contract['blocks'][0] if contract['blocks'] else {'id': 'p1', 'text': '未识别'}
    return client.post(f'/legal/contracts/{cid}/reviews', json={
        'our_role': '采购方', 'contract_type': '采购合同', 'our_party': {'block_id': block['id'], 'quote': block['text'][:2]},
        'model_id': 'claude-test', 'external_processing_provider': 'anthropic',
        'external_processing_confirmed': True, 'review_tier': tier})


def complete(client, contract, block_id, replacement, *, decision='draft'):
    response = start(client, contract)
    assert response.status_code == 202, response.text
    rid = response.json()['id']
    review = store.review(rid)
    block = next(b for b in contract['blocks'] if b['id'] == block_id)
    p = review['payload']
    p['findings'] = [{'id': 'f1', 'block_id': block_id, 'original_quote': block['text'], 'title': '合成修改建议',
                      'kind': 'commercial', 'severity': 'medium', 'impact': '调整合同约定。', 'reason': '合成测试。',
                      'suggested_text': replacement, 'revision_allowed': True, 'verification_status': 'uncertain',
                      'evidence_status': 'not_applicable', 'missing_facts': [], 'citations': [], 'policy_ids': []}]
    token = store.claim(rid)
    store.checkpoint(rid, token, p, 'completed')
    response = client.patch(f'/legal/reviews/{rid}/findings/f1', json={'decision': decision, 'text': replacement, 'expected_version': 0})
    assert response.status_code == 200, response.text
    return rid


def test_private_originals_table_redaction_and_assets_through_api(api):
    client, actor, path = api
    raw = word()
    contract = upload(client, raw, 'docx')
    cid = contract['id']
    encoded = json.dumps(contract, ensure_ascii=False)
    assert all(value not in encoded for value in ('张三', 'PRIVATE ALT', 'private@example.test', 'source_ref', 'source_part', 'original_b64'))
    assert contract['original_available'] and contract['reviewable']
    assert contract['document']['body'][1]['rows'][2]['cells'][0]['col_span'] == 2
    assert '张三'.encode() not in path.read_bytes()
    assert client.get(f'/legal/contracts/{cid}/original').content == raw
    original_text = client.get(f'/legal/contracts/{cid}/original-text')
    assert '张三' in original_text.text and original_text.headers['cache-control'] == 'no-store'
    asset_id = contract['assets'][0]['id']
    picture = client.get(f'/legal/contracts/{cid}/assets/{asset_id}')
    assert picture.status_code == 200 and picture.headers['content-type'] == 'image/png'
    assert Image.open(io.BytesIO(picture.content)).size == (360, 180)
    assert picture.headers['cache-control'] == 'no-store'
    assert client.get(f'/legal/contracts/{cid}/assets/missing').status_code == 404
    actor['matter_role'] = 'viewer'
    assert client.get(f'/legal/contracts/{cid}').json()['original_available'] is False
    assert client.get(f'/legal/contracts/{cid}/original').status_code == 403
    assert client.get(f'/legal/contracts/{cid}/assets/{asset_id}').status_code == 403
    actor['id'] = 'foreign-tenant'
    assert client.get(f'/legal/contracts/{cid}').status_code == 403


def test_structured_word_exports_with_one_user_confirmation_and_no_model_veto(api):
    client, _, _ = api
    raw = word()
    contract = upload(client, raw, 'docx')
    block = contract['blocks'][0]
    rid = complete(client, contract, block['id'], block['text'].replace('90', '30'))
    state = client.get(f'/legal/reviews/{rid}').json()['export_state']
    assert state['pending_count'] == 1 and state['parsing_issue_count'] > 0
    result = client.get(f'/legal/reviews/{rid}/export', params={'format': 'docx'})
    assert result.status_code == 409 and result.json()['detail']['pending_count'] == 1
    result = client.get(f'/legal/reviews/{rid}/export', params={'format': 'docx', 'confirmed': 'true', 'fingerprint': state['fingerprint']})
    assert result.status_code == 200 and result.headers['x-legal-appendix-count'] == '0'
    with zipfile.ZipFile(io.BytesIO(raw)) as old, zipfile.ZipFile(io.BytesIO(result.content)) as new:
        media = next(name for name in old.namelist() if name.startswith('word/media/'))
        assert old.read(media) == new.read(media)
        assert old.read('word/header1.xml') == new.read('word/header1.xml')
        root = etree.fromstring(new.read('word/document.xml'))
        assert root.xpath('//w:ins//w:t/text()', namespaces=NS) == ['3']
        assert root.xpath('//w:del//w:delText/text()', namespaces=NS) == ['9']
        assert root.xpath('//w:gridSpan/@w:val', namespaces=NS) == ['2']


def test_non_native_word_change_is_an_explicit_appendix_and_still_exports_directly(api):
    client, _, _ = api
    raw = word()
    contract = upload(client, raw, 'docx')
    header = next(block for block in contract['blocks'] if block['kind'] == 'header')
    replacement = '通知应通过双方约定的电子邮件送达。'
    rid = complete(client, contract, header['id'], replacement, decision='accepted')
    state = client.get(f'/legal/reviews/{rid}').json()['export_state']
    assert state['ready'] and state['pending_count'] == 0 and state['appendix_count'] == 1
    response = client.get(f'/legal/reviews/{rid}/export', params={'format': 'docx'})
    assert response.status_code == 200 and response.headers['x-legal-appendix-count'] == '1'
    with zipfile.ZipFile(io.BytesIO(raw)) as old, zipfile.ZipFile(io.BytesIO(response.content)) as new:
        assert old.read('word/header1.xml') == new.read('word/header1.xml')
        assert replacement in new.read('word/document.xml').decode()


def test_pdf_keeps_source_pixels_and_text_with_annotations_and_complete_appendix(api):
    client, _, _ = api
    raw = pdf()
    contract = upload(client, raw, 'pdf')
    assert contract['assets'] and any(b.get('bbox') for b in contract['blocks'])
    block = next(block for block in contract['blocks'] if '90 days' in block['text'])
    rid = complete(client, contract, block['id'], block['text'].replace('90 days', '30 days'), decision='accepted')
    response = client.get(f'/legal/reviews/{rid}/export', params={'format': 'pdf'})
    assert response.status_code == 200, response.text
    source, edited = PdfReader(io.BytesIO(raw)), PdfReader(io.BytesIO(response.content))
    assert len(edited.pages) == len(source.pages) + 1
    assert '90 days' in edited.pages[0].extract_text() and '30 days' not in edited.pages[0].extract_text()
    assert '30 days' in edited.pages[-1].extract_text()
    comments = [annotation.get_object().get('/Contents', '') for annotation in edited.pages[0]['/Annots']]
    assert any('30 days' in comment for comment in comments)
    before_image = list(source.pages[0].images)[0].data
    assert list(edited.pages[0].images)[0].data == before_image
    assert client.get(f'/legal/contracts/{contract["id"]}/original').content == raw


def test_scan_only_is_viewable_but_cannot_start_an_empty_model_review(api):
    client, _, _ = api
    raw = pdf(scan_only=True)
    contract = upload(client, raw, 'pdf')
    assert contract['blocks'] == [] and contract['parsing_status'] == 'needs_ocr'
    assert not contract['reviewable'] and contract['original_available']
    assert contract['document']['issues'] and client.get(f'/legal/contracts/{contract["id"]}/original').content == raw
    response = start(client, contract)
    assert response.status_code == 422 and '尚未识别' in response.text


def test_model_context_keeps_cells_and_numbering_without_source_or_pixel_leaks(monkeypatch):
    monkeypatch.delenv('LEGAL_DOCUMENT_PARSER_URL', raising=False)
    monkeypatch.delenv('LEGAL_DOCUMENT_PARSER_TOKEN', raising=False)
    parsed = documents.parse_contract(word(), 'synthetic.docx')
    masked, _ = documents.redact_blocks(parsed['blocks'])
    model_blocks = context.model_blocks(masked)
    structure = context.model_structure(parsed['document'])
    encoded = json.dumps({'blocks': context.model_request_blocks(model_blocks, structure), 'structure': structure}, ensure_ascii=False)
    assert all(value not in encoded for value in ('张三', 'PRIVATE ALT', 'source_part', 'source_ref', 'redaction_spans', 'anchor', 'data_b64'))
    table = structure['body'][1]
    assert table['rows'][2][0]['span'] == [1, 2]
    prompt = review_chat.document({}, model_blocks, [], [], [], parsed['warnings'], parsed['document'])
    assert '<table id=' in prompt and 'colspan="2"' in prompt and '联系人' in prompt
    assert '张三' not in prompt and '未向模型发送图像' in prompt
    for block in masked:
        assert prompt.count('[' + block['id'] + ']') == 1


def test_repeated_cell_values_require_unambiguous_source_identity():
    cells = [{'id': f'p{i}', 'text': '30天', 'table': {'id': 't1', 'cell_id': f'c{i}', 'row': i, 'col': 1}}
             for i in (1, 2)]
    assert review_chat.anchor('p2', '30天', '60天', cells) == ('p2', '30天')
    assert review_chat.anchor(None, '30天', '60天', cells) == (None, '')
    assert review_chat.anchor('p1', '90天', '60天', cells) == (None, '')
    assert review_chat.anchor(None, '', '30天改为60天', cells) == (None, '')
    assert review_chat.anchor('p0', '30天', '60天', [{'id': 'p0', 'text': '普通签约条款'}] + cells) == (None, '')


@pytest.mark.parametrize('ocr_table', [False, True])
def test_word_ocr_is_once_per_asset_but_each_image_placement_has_its_own_anchor(monkeypatch, ocr_table):
    from legal.document_enrichment import enrich_docx_images
    from legal.docx_structure import parse_docx
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '联系人：'
    table.cell(0, 1).paragraphs[0].add_run().add_picture(io.BytesIO(png()))
    doc.add_picture(io.BytesIO(png()))
    stream = io.BytesIO()
    doc.save(stream)
    data = stream.getvalue()
    parsed = parse_docx(data)
    calls = []
    monkeypatch.setattr(document_ocr, 'ocr_status', lambda: {'configured': True})

    def recognize(raw, media, *, budget):
        calls.append(budget)
        item = {'text': '张三', 'bbox': [0, 0, 180, 80]}
        if ocr_table:
            item['table'] = {'id': 'ocr-table1', 'cell_id': 'ocr-cell1', 'row': 0, 'col': 0, 'row_span': 1, 'col_span': 1}
        return {'status': 'complete', 'blocks': [item]}

    monkeypatch.setattr(document_ocr, 'recognize_image', recognize)
    enrich_docx_images(parsed, data)
    ocr = [block for block in parsed['blocks'] if block['kind'] == 'ocr']
    assert len(calls) == 1 and len(ocr) == 2 and ocr[0]['id'] != ocr[1]['id']
    assert (ocr[0]['ancestor_tables'][0] if ocr_table else ocr[0]['table'])['col'] == 1
    redacted, mapping = documents.redact_blocks(parsed['blocks'])
    assert '张三' in mapping.values()
    assert all('张三' not in block['text'] for block in redacted)
    assert all(block['anchor'] is None and not block['native_editable'] for block in ocr)


@pytest.mark.parametrize('source_format', ['PNG', 'JPEG'])
def test_raster_previews_strip_source_metadata(source_format):
    from legal.document_assets import raster_preview
    source = Image.new('RGB', (240, 120), '#426a94')
    exif = source.getexif()
    exif[0x010E] = 'PRIVATE document description'
    data = io.BytesIO()
    source.save(data, format=source_format, icc_profile=b'PRIVATE embedded profile', exif=exif)
    preview, mime = raster_preview(data.getvalue())
    assert mime == 'image/png'
    with Image.open(io.BytesIO(preview)) as result:
        assert result.size == source.size
        assert not result.info and not result.getexif()
    assert b'PRIVATE' not in preview


def test_structure_changes_invalidate_a_previously_confirmed_export(api):
    client, _, _ = api
    contract = upload(client, word(), 'docx')
    block = contract['blocks'][0]
    rid = complete(client, contract, block['id'], block['text'].replace('90', '30'))
    state = client.get(f'/legal/reviews/{rid}').json()['export_state']
    with store.transaction() as conn:
        saved = store.decode(conn.execute('SELECT payload FROM legal_contracts WHERE id=?', (contract['id'],)).fetchone()[0])
        saved['document']['body'][1]['rows'][0]['cells'][0]['col_span'] = 2
        conn.execute('UPDATE legal_contracts SET payload=? WHERE id=?', (store.encode(saved), contract['id']))
    response = client.get(f'/legal/reviews/{rid}/export', params={'format': 'docx', 'confirmed': 'true', 'fingerprint': state['fingerprint']})
    assert response.status_code == 409 and response.json()['detail']['code'] == 'export_stale'


def test_version_requires_the_structured_backend(api):
    assert api[0].get('/legal/version').json()['document_structure_version'] == 1

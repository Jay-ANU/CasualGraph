"""Native structure, redaction and export fidelity using small synthetic files."""
from copy import deepcopy
import io
import json
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image
import pytest

from legal.docx_structure import parse_docx
from legal.docx_redlines import NS, W, UnsupportedRevision, paragraph_text, redline_docx, redline_docx_result, visible_paragraph_text
from legal.redaction_detection import redact


def package(doc):
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def xml(data, part='word/document.xml'):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return etree.fromstring(archive.read(part))


def picture():
    out = io.BytesIO()
    Image.new('RGB', (24, 16), 'navy').save(out, format='PNG')
    out.seek(0)
    return out


def view(data, accept):
    root = xml(data)
    for tag in ('ins', 'del'):
        for revision in list(root.xpath(f'//w:{tag}', namespaces=NS)):
            parent = revision.getparent()
            at = parent.index(revision)
            if (accept and tag == 'ins') or (not accept and tag == 'del'):
                for child in list(revision):
                    for text in child.xpath('.//w:delText', namespaces=NS):
                        text.tag = qn('w:t')
                    parent.insert(at, child)
                    at += 1
            parent.remove(revision)
    return root


def table_block(bid, text, row, col, *, table_id='table1', header=False, col_span=1, row_span=1):
    return {'id': bid, 'text': text, 'table': {'id': table_id, 'cell_id': f'{table_id}-{row}-{col}',
            'row': row, 'col': col, 'row_span': row_span, 'col_span': col_span, 'header': header}}


def assert_mask_roundtrip(originals, masked, mapping):
    for source, result in zip(originals, masked):
        restored = result['text']
        for token, value in mapping.items():
            restored = restored.replace(token, value)
        assert restored == source['text']
        for span in result['redaction_spans']:
            assert result['text'][span['start']:span['end']] == span['token']
            assert source['text'][span['original_start']:span['original_end']] == mapping[span['token']]


def test_ordered_tables_merge_coordinates_and_original_paragraph_anchors():
    doc = Document()
    doc.add_paragraph('合同正文前段。')
    table = doc.add_table(rows=4, cols=3)
    for r in range(4):
        for c in range(3):
            table.cell(r, c).text = f'行{r}列{c}'
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(1, 0).merge(table.cell(1, 1))
    table.cell(2, 2).merge(table.cell(3, 2))
    table.cell(3, 1).text = ''
    doc.add_paragraph('合同正文后段。')
    raw = package(doc)
    result = parse_docx(raw)
    assert [node['type'] for node in result['document']['body']] == ['paragraph', 'table', 'paragraph']
    structured = result['document']['body'][1]
    assert [cell['col'] for cell in structured['rows'][1]['cells']] == [0, 2]
    assert structured['rows'][1]['cells'][0]['col_span'] == 2
    assert structured['rows'][2]['cells'][2]['row_span'] == 2
    assert len(structured['rows'][3]['cells']) == 2
    assert structured['rows'][3]['cells'][1]['children'] == []
    assert all(cell['header'] for cell in structured['rows'][0]['cells'])
    paragraphs = xml(raw).xpath('//w:body//w:p[not(ancestor::w:txbxContent)]', namespaces=NS)
    for block in result['blocks']:
        assert block['id'] == f"p{block['anchor'] + 1}"
        assert paragraph_text(paragraphs[block['anchor']]).strip() == block['text']
        assert block['native_editable'] is True
        if block['text'] == '行3列2':
            assert block['table']['row_span'] == 2
            assert block['table']['row'] == 2
    assert redline_docx(raw, result['blocks'], {}) == raw


def test_nested_table_keeps_cell_context_and_sibling_order():
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    cell = outer.cell(0, 0)
    cell.text = '外层前文。'
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = '内层条款。'
    cell.add_paragraph('外层后文。')
    result = parse_docx(package(doc))
    children = result['document']['body'][0]['rows'][0]['cells'][0]['children']
    assert [node['type'] for node in children] == ['paragraph', 'table', 'paragraph']
    by_text = {block['text']: block for block in result['blocks']}
    assert by_text['外层前文。']['table']['id'] == by_text['外层后文。']['table']['id']
    assert by_text['内层条款。']['table']['id'] != by_text['外层前文。']['table']['id']


def test_legacy_horizontal_and_vertical_merge_share_one_logical_cell():
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    for row in range(2):
        for col in range(3):
            cell = table.cell(row, col)
            cell.text = f'内容{row}{col}'
            if col < 2:
                props = cell._tc.get_or_add_tcPr()
                hm = OxmlElement('w:hMerge'); hm.set(qn('w:val'), 'restart' if col == 0 else 'continue'); props.append(hm)
                vm = OxmlElement('w:vMerge'); vm.set(qn('w:val'), 'restart' if row == 0 else 'continue'); props.append(vm)
    result = parse_docx(package(doc))
    rows = result['document']['body'][0]['rows']
    merged = rows[0]['cells'][0]
    assert (merged['row_span'], merged['col_span']) == (2, 2)
    assert len(merged['children']) == 4
    assert len(rows[1]['cells']) == 1 and rows[1]['cells'][0]['col'] == 2
    selected = [block for block in result['blocks'] if block['text'][-1] in ('0', '1')]
    assert len({block['table']['cell_id'] for block in selected}) == 1
    assert all(block['table']['row_span'] == 2 and block['table']['col_span'] == 2 for block in selected)


def test_malformed_legacy_merge_discloses_problem_without_losing_cell_text():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    for col in range(2):
        table.cell(0, col).text = f'应保留{col}'
    table.cell(0, 0)._tc.get_or_add_tcPr().append(OxmlElement('w:hMerge'))
    result = parse_docx(package(doc))
    assert [block['text'] for block in result['blocks']] == ['应保留0', '应保留1']
    assert len(result['document']['body'][0]['rows'][0]['cells']) == 2
    assert any(issue['code'] == 'table_merge_unresolved' for issue in result['document']['issues'])


def test_images_are_referenced_once_without_alt_text_in_metadata():
    doc = Document()
    host = doc.add_paragraph('图示对应本段。')
    drawing = host.add_run().add_picture(picture())
    drawing._inline.docPr.set('descr', 'PRIVATE ALT: 联系人张三 legal@example.test')
    doc.add_picture(picture())
    result = parse_docx(package(doc))
    assert len(result['assets']) == 1
    assert result['assets'][0]['media_type'] == 'image/png'
    assert (result['assets'][0]['width'], result['assets'][0]['height']) == (24, 16)
    images = [node for node in result['document']['body'] if node['type'] == 'image']
    assert len(images) == 2 and images[0]['asset_id'] == images[1]['asset_id']
    assert images[0]['host_block_id'] == result['blocks'][0]['id']
    assert all('block_ids' not in image for image in images)
    assert 'PRIVATE ALT' not in json.dumps(result, ensure_ascii=False)
    assert all(block['text'] != '' for block in result['blocks'])


def add_textbox(paragraph, value):
    pict = OxmlElement('w:pict')
    shape = etree.SubElement(pict, '{urn:schemas-microsoft-com:vml}shape')
    box = etree.SubElement(shape, '{urn:schemas-microsoft-com:vml}textbox')
    content = etree.SubElement(box, qn('w:txbxContent'))
    p = etree.SubElement(content, qn('w:p'))
    run = etree.SubElement(p, qn('w:r'))
    etree.SubElement(run, qn('w:t')).text = value
    paragraph.add_run()._r.append(pict)


def test_readonly_sections_include_textboxes_headers_footers_once():
    doc = Document()
    host = doc.add_paragraph('正文前段。')
    add_textbox(host, '联系人：张三')
    doc.add_paragraph('正文后段。')
    doc.sections[0].header.paragraphs[0].text = '页眉主体信息。'
    doc.sections[0].footer.paragraphs[0].text = '页脚通知说明。'
    raw = package(doc)
    result = parse_docx(raw)
    assert {section['kind'] for section in result['document']['sections']} == {'header', 'footer', 'textbox'}
    assert [block['id'] for block in result['blocks'] if block['anchor'] is not None] == ['p1', 'p2']
    assert sum(block['text'] == '联系人：张三' for block in result['blocks']) == 1
    assert all(block['native_editable'] is False and block['source_ref'] for block in result['blocks'] if block['anchor'] is None)
    header = next(block for block in result['blocks'] if block['kind'] == 'header')
    exported = redline_docx_result(raw, result['blocks'], {header['id']: '建议的页眉说明。'})
    assert exported['applied_block_ids'] == []
    assert exported['unapplied'][0]['block_id'] == header['id']
    with zipfile.ZipFile(io.BytesIO(raw)) as old, zipfile.ZipFile(io.BytesIO(exported['content'])) as new:
        assert old.read(header['source_ref']['part']) == new.read(header['source_ref']['part'])
    assert '建议的页眉说明。' in ''.join(xml(exported['content']).xpath('//w:t/text()', namespaces=NS))


def test_textbox_only_document_has_reviewable_text_without_a_fake_body_anchor():
    doc = Document()
    add_textbox(doc.add_paragraph(), '本合同的全部文字位于文本框。')
    parsed = parse_docx(package(doc))
    assert len(parsed['blocks']) == 1
    assert parsed['blocks'][0]['text'] == '本合同的全部文字位于文本框。'
    assert parsed['blocks'][0]['anchor'] is None and parsed['blocks'][0]['native_editable'] is False
    assert parsed['document']['body'] == []
    assert parsed['document']['sections'][0]['body'][0]['block_id'] == parsed['blocks'][0]['id']


def test_textbox_image_keeps_private_field_context_for_ocr(monkeypatch):
    from docx.text.paragraph import Paragraph
    from legal import document_ocr
    from legal.document_enrichment import enrich_docx_images

    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '联系人：'
    host = table.cell(0, 1).paragraphs[0]
    add_textbox(host, '')
    textbox = host._p.xpath('.//w:txbxContent/w:p')[0]
    pixels = io.BytesIO()
    Image.new('RGB', (400, 200), 'white').save(pixels, format='PNG')
    pixels.seek(0)
    Paragraph(textbox, host._parent).add_run().add_picture(pixels)
    raw = package(doc)
    parsed = parse_docx(raw)
    image = parsed['document']['sections'][0]['body'][0]
    assert image['type'] == 'image' and len(image['ancestor_tables']) == 1
    context = image['ancestor_tables'][0]
    assert context['id'] == parsed['document']['body'][0]['id']
    assert (context['row'], context['col']) == (0, 1)
    monkeypatch.setattr(document_ocr, 'ocr_status', lambda: {'configured': True})
    monkeypatch.setattr(document_ocr, 'recognize_image', lambda *a, **kw: {'blocks': [
        {'text': '张三', 'table': {'id': 'inner1', 'cell_id': 'cell1', 'row': 0, 'col': 0,
                                'row_span': 1, 'col_span': 1}}]})
    enrich_docx_images(parsed, raw)
    masked, mapping = redact(parsed['blocks'])
    assert set(mapping.values()) == {'张三'}
    assert next(b for b in masked if b['kind'] == 'ocr')['text'] == next(iter(mapping))
    assert_mask_roundtrip(parsed['blocks'], masked, mapping)


def numbering(doc, *, sensitive_label=False):
    root = doc.part.numbering_part.element
    abstract = OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'), '900')
    for level, fmt, pattern in ((0, 'decimal', 'legal@example.test %1.' if sensitive_label else '%1.'), (1, 'lowerLetter', '%1.%2)')):
        node = OxmlElement('w:lvl'); node.set(qn('w:ilvl'), str(level))
        for name, value in (('start', '1'), ('numFmt', fmt), ('lvlText', pattern)):
            child = OxmlElement('w:' + name); child.set(qn('w:val'), value); node.append(child)
        abstract.append(node)
    root.append(abstract)
    num = OxmlElement('w:num'); num.set(qn('w:numId'), '900')
    ref = OxmlElement('w:abstractNumId'); ref.set(qn('w:val'), '900'); num.append(ref)
    override = OxmlElement('w:lvlOverride'); override.set(qn('w:ilvl'), '0')
    start = OxmlElement('w:startOverride'); start.set(qn('w:val'), '5'); override.append(start); num.append(override)
    root.append(num)


def numbered(doc, level, value):
    paragraph = doc.add_paragraph(value)
    props = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    props.get_or_add_numId().val = 900
    props.get_or_add_ilvl().val = level


def test_numbering_levels_restarts_and_overrides_are_display_metadata():
    doc = Document()
    numbering(doc)
    for level in (0, 1, 1, 0, 1):
        numbered(doc, level, '约定内容。')
    result = parse_docx(package(doc))
    assert [block['label'] for block in result['blocks']] == ['5.', '5.a)', '5.b)', '6.', '6.a)']
    assert all(block['text'] == '约定内容。' for block in result['blocks'])


def test_style_numbering_and_headings_are_resolved_without_style_names():
    doc = Document()
    doc.add_heading('标题文本', level=2)
    doc.add_paragraph('首项。', style='List Number')
    doc.add_paragraph('次项。', style='List Number')
    result = parse_docx(package(doc))
    assert result['blocks'][0]['heading_level'] == 2
    assert [block['label'] for block in result['blocks'][1:]] == ['1.', '2.']
    assert 'ListNumber' not in json.dumps(result)


def test_numbering_label_uses_the_same_redaction_namespace_as_text():
    doc = Document()
    numbering(doc, sensitive_label=True)
    numbered(doc, 0, '通知邮箱 legal@example.test。')
    result = parse_docx(package(doc))
    masked, mapping = redact(result['blocks'])
    assert list(mapping.values()) == ['legal@example.test']
    token = next(iter(mapping))
    assert token in masked[0]['label'] and token in masked[0]['text']
    assert 'legal@example.test' not in json.dumps(masked, ensure_ascii=False)
    assert masked[0]['label_redaction_spans'][0]['token'] == token
    assert_mask_roundtrip(result['blocks'], masked, mapping)


def test_adjacent_and_merged_field_cells_mask_values_keep_duties_and_amounts():
    blocks = [table_block('p1', '联系人：', 0, 0), table_block('p2', '张三负责验收。', 0, 1),
              table_block('p3', '付款期限', 1, 0), table_block('p4', '30天', 1, 1),
              table_block('p5', '联系地址', 2, 0, col_span=2), table_block('p6', '上海市测试路100号', 2, 2)]
    masked, mapping = redact(blocks)
    assert '张三' in mapping.values() and '上海市测试路100号' in mapping.values()
    assert '负责验收。' in masked[1]['text']
    assert masked[0]['text'] == '联系人：' and masked[3]['text'] == '30天'
    assert_mask_roundtrip(blocks, masked, mapping)
    unmasked, remaining = redact(blocks, excluded_terms=['张三'])
    assert '张三' in unmasked[1]['text'] and '张三' not in remaining.values()


def test_real_docx_contact_table_preserves_context_through_redaction():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '联系人：'
    table.cell(0, 1).text = '张三'
    parsed = parse_docx(package(doc))
    masked, mapping = redact(parsed['blocks'])
    assert list(mapping.values()) == ['张三']
    assert masked[0]['text'] == '联系人：' and masked[1]['text'] == next(iter(mapping))
    assert_mask_roundtrip(parsed['blocks'], masked, mapping)


def test_nested_docx_value_retains_every_enclosing_field_cell():
    doc = Document()
    outer = doc.add_table(rows=2, cols=2)
    outer.cell(0, 0).text = '联系人：'
    inner = outer.cell(0, 1).add_table(rows=1, cols=1)
    innermost = inner.cell(0, 0).add_table(rows=1, cols=1)
    innermost.cell(0, 0).text = '张三'
    outer.cell(1, 0).text = '金额'
    outer.cell(1, 1).add_table(rows=1, cols=1).cell(0, 0).text = '10000'
    parsed = parse_docx(package(doc))
    name = next(b for b in parsed['blocks'] if b['text'] == '张三')
    assert len(name['ancestor_tables']) == 2
    assert name['ancestor_tables'][0]['id'] == parsed['document']['body'][0]['id']
    assert (name['ancestor_tables'][0]['row'], name['ancestor_tables'][0]['col']) == (0, 1)
    assert len({context['id'] for context in [*name['ancestor_tables'], name['table']]}) == 3
    masked, mapping = redact(parsed['blocks'])
    assert set(mapping.values()) == {'张三'}
    assert next(b for b in masked if b['id'] == name['id'])['text'] == next(iter(mapping))
    assert '10000' in [b['text'] for b in masked]
    assert_mask_roundtrip(parsed['blocks'], masked, mapping)


@pytest.mark.parametrize('header', [False, True])
def test_ocr_inner_table_uses_outer_field_without_crossing_other_cells(header):
    label = table_block('p1', '姓名' if header else '联系人：', 0, 0, header=header)
    row, col = (1, 0) if header else (0, 1)
    name = {**table_block('ocr-image1-1', '张三', 0, 0, table_id='image1-table1'),
            'kind': 'ocr', 'native_editable': False, 'asset_id': 'asset1',
            'ancestor_tables': [table_block('unused', '', row, col)['table'],
                                table_block('unused', '', 2, 1, table_id='nested1')['table']]}
    unrelated = {**table_block('ocr-image2-1', '李四', 0, 0, table_id='image2-table1'),
                 'kind': 'ocr', 'asset_id': 'asset2'}
    amount_label = table_block('p2', '金额', 0 if header else 2, 1 if header else 0, header=header)
    amount = {**table_block('ocr-image3-1', '10000', 0, 0, table_id='image3-table1'),
              'kind': 'ocr', 'ancestor_tables': [table_block('unused', '', 1 if header else 2, 1)['table']]}
    blocks = [label, name, unrelated, amount_label, amount]
    masked, mapping = redact(blocks)
    assert set(mapping.values()) == {'张三'}
    assert masked[1]['text'] == next(iter(mapping))
    assert masked[2]['text'] == '李四' and masked[4]['text'] == '10000'
    assert_mask_roundtrip(blocks, masked, mapping)
    unmasked, remaining = redact(blocks, excluded_terms=['张三'])
    assert unmasked[1]['text'] == '张三' and remaining == {}


@pytest.mark.parametrize('explicit', [False, True])
def test_column_headers_mask_person_and_address_cells_without_masking_amounts(explicit):
    blocks = [table_block('p1', '姓名', 0, 0, header=explicit), table_block('p2', '金额', 0, 1, header=explicit),
              table_block('p3', '联系地址', 0, 2, header=explicit), table_block('p4', '张三', 1, 0),
              table_block('p5', '10000', 1, 1), table_block('p6', '北京市测试路1号', 1, 2),
              table_block('p7', '李四', 2, 0), table_block('p8', '20000', 2, 1), table_block('p9', '上海市测试路2号', 2, 2)]
    masked, mapping = redact(blocks)
    assert {'张三', '李四', '北京市测试路1号', '上海市测试路2号'}.issubset(mapping.values())
    assert masked[4]['text'] == '10000' and masked[7]['text'] == '20000'
    assert_mask_roundtrip(blocks, masked, mapping)


def test_blank_adjacent_cell_does_not_shift_field_to_unrelated_column_or_table():
    blocks = [table_block('p1', '联系人', 0, 0), table_block('p2', '交付条款', 0, 2),
              table_block('p3', '相邻说明', 0, 1, table_id='table2')]
    masked, mapping = redact(blocks)
    assert mapping == {} and [b['text'] for b in masked] == [b['text'] for b in blocks]


def test_mixed_picture_text_revisions_preserve_object_and_accept_reject_semantics():
    doc = Document()
    paragraph = doc.add_paragraph()
    run = paragraph.add_run('付款90')
    run.bold = True
    run.add_picture(picture())
    run.add_text('天。')
    raw = package(doc)
    result = parse_docx(raw)
    assert result['blocks'][0]['native_editable'] is True
    revised = redline_docx(raw, result['blocks'], {'p1': '付款120天。'})
    before_object = xml(raw).xpath('//w:drawing', namespaces=NS)[0]
    after_object = xml(revised).xpath('//w:drawing', namespaces=NS)[0]
    assert etree.tostring(before_object) == etree.tostring(after_object)
    assert visible_paragraph_text(view(revised, False).xpath('//w:body/w:p', namespaces=NS)[0]) == '付款90天。'
    assert visible_paragraph_text(view(revised, True).xpath('//w:body/w:p', namespaces=NS)[0]) == '付款120天。'
    for accept, expected in ((False, '付款90'), (True, '付款120')):
        paragraph = view(revised, accept).xpath('//w:body/w:p', namespaces=NS)[0]
        preceding = ''
        for run in paragraph:
            if run.xpath('.//w:drawing', namespaces=NS):
                break
            preceding += ''.join(run.xpath('.//w:t/text()', namespaces=NS))
        assert preceding == expected
    with zipfile.ZipFile(io.BytesIO(raw)) as old, zipfile.ZipFile(io.BytesIO(revised)) as new:
        for name in old.namelist():
            if name.startswith('word/media/'):
                assert old.read(name) == new.read(name)


def test_crossing_image_anchor_keeps_original_and_appends_exact_choice():
    doc = Document()
    paragraph = doc.add_paragraph('ABC')
    paragraph.add_run().add_picture(picture())
    paragraph.add_run('DEF')
    raw = package(doc)
    parsed = parse_docx(raw)
    with pytest.raises(UnsupportedRevision):
        redline_docx(raw, parsed['blocks'], {'p1': 'XYZ'})
    result = redline_docx_result(raw, parsed['blocks'], {'p1': 'XYZ'})
    assert result['applied_block_ids'] == [] and result['unapplied'][0]['block_id'] == 'p1'
    first = xml(result['content']).xpath('//w:body/w:p', namespaces=NS)[0]
    assert etree.tostring(first) == etree.tostring(xml(raw).xpath('//w:body/w:p', namespaces=NS)[0])
    texts = xml(result['content']).xpath('//w:t/text()', namespaces=NS)
    assert '未写入原位置的修改建议' in texts and 'XYZ' in texts
    assert any('没有替换原位置' in text for text in texts)


def test_unsupported_paragraph_appendix_preserves_original_and_other_native_changes():
    doc = Document()
    paragraph = doc.add_paragraph('原始复杂条款。')
    bookmark = OxmlElement('w:bookmarkStart'); bookmark.set(qn('w:id'), '100'); bookmark.set(qn('w:name'), 'test')
    paragraph._p.insert(0, bookmark)
    doc.add_paragraph('付款90天。')
    raw = package(doc)
    parsed = parse_docx(raw)
    assert parsed['blocks'][0]['native_editable'] is False
    result = redline_docx_result(raw, parsed['blocks'], {'p1': '完整复杂条款修改建议。', 'p2': '付款30天。'})
    assert result['applied_block_ids'] == ['p2']
    assert [item['block_id'] for item in result['unapplied']] == ['p1']
    root = xml(result['content'])
    assert etree.tostring(root.xpath('//w:body/w:p', namespaces=NS)[0]) == etree.tostring(xml(raw).xpath('//w:body/w:p', namespaces=NS)[0])
    assert root.xpath('//w:body/w:p[2]//w:ins', namespaces=NS)
    assert '完整复杂条款修改建议。' in root.xpath('//w:t/text()', namespaces=NS)
    assert root.find(qn('w:body'))[-1].tag == qn('w:sectPr')


def test_content_control_is_kept_intact_and_its_edit_is_explicitly_unapplied():
    doc = Document()
    paragraph = doc.add_paragraph('控件内条款。')
    body = paragraph._p.getparent()
    index = body.index(paragraph._p)
    control = OxmlElement('w:sdt')
    props = OxmlElement('w:sdtPr'); props.append(OxmlElement('w:dataBinding')); control.append(props)
    content = OxmlElement('w:sdtContent'); control.append(content)
    content.append(paragraph._p)
    body.insert(index, control)
    raw = package(doc)
    parsed = parse_docx(raw)
    assert parsed['blocks'][0]['native_editable'] is False
    result = redline_docx_result(raw, parsed['blocks'], {'p1': '完整新条款。'})
    assert result['unapplied'][0]['block_id'] == 'p1'
    assert etree.tostring(xml(raw).xpath('//w:sdt', namespaces=NS)[0]) == etree.tostring(xml(result['content']).xpath('//w:sdt', namespaces=NS)[0])


@pytest.mark.parametrize('failure', ['unknown', 'stale', 'duplicate', 'bad_anchor', 'unsafe_text'])
def test_appendix_never_downgrades_integrity_failures(failure):
    doc = Document(); doc.add_paragraph('原始条款。')
    raw = package(doc)
    originals = parse_docx(raw)['blocks']
    replacements = {'p1': '建议条款。'}
    if failure == 'unknown':
        replacements = {'missing': '建议条款。'}
    elif failure == 'stale':
        originals[0]['text'] = '旧版原文。'
    elif failure == 'duplicate':
        originals.append(deepcopy(originals[0]))
    elif failure == 'bad_anchor':
        originals[0]['anchor'] = 9999
    else:
        replacements['p1'] += '\x00'
    with pytest.raises(ValueError):
        redline_docx_result(raw, originals, replacements)


def test_stale_ancillary_source_cannot_be_hidden_in_appendix():
    doc = Document(); doc.add_paragraph('正文。')
    doc.sections[0].header.paragraphs[0].text = '原页眉。'
    raw = package(doc)
    originals = parse_docx(raw)['blocks']
    header = next(block for block in originals if block['kind'] == 'header')
    header['text'] = '过期页眉。'
    with pytest.raises(ValueError, match='已变化'):
        redline_docx_result(raw, originals, {header['id']: '修改页眉。'})


def test_native_source_reference_must_match_the_actual_paragraph():
    doc = Document(); doc.add_paragraph('相同内容。'); doc.add_paragraph('相同内容。')
    raw = package(doc)
    originals = parse_docx(raw)['blocks']
    originals[0]['source_ref']['paragraph'] = 1
    with pytest.raises(ValueError, match='定位'):
        redline_docx_result(raw, originals, {'p1': '新的内容。'})


def test_two_ancillary_ids_cannot_replace_the_same_source_twice():
    doc = Document(); doc.add_paragraph('正文。')
    doc.sections[0].header.paragraphs[0].text = '页眉。'
    raw = package(doc)
    originals = parse_docx(raw)['blocks']
    header = next(block for block in originals if block['kind'] == 'header')
    alias = {**deepcopy(header), 'id': 'auxp-alias'}
    originals.append(alias)
    with pytest.raises(ValueError, match='重复'):
        redline_docx_result(raw, originals, {header['id']: '建议一。', alias['id']: '建议二。'})

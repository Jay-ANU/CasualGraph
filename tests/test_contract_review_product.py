"""Contract product regression tests. Providers are mocked; no model credit is used."""
import io
import json
import sqlite3
import time
import zipfile

import httpx
import pytest
from cryptography.fernet import Fernet
from docx import Document
from fastapi import HTTPException

from legal import contract_documents as docs, external_law as laws, review_engine as engine, review_store as store


@pytest.fixture
def db(tmp_path, monkeypatch):
    path = tmp_path / 'legal.db'
    def connect():
        c = sqlite3.connect(path)
        c.row_factory = sqlite3.Row
        c.execute('CREATE TABLE IF NOT EXISTS org_members (org_id TEXT,user_id TEXT,role TEXT)')
        c.execute('''CREATE TABLE IF NOT EXISTS audit_events (id INTEGER PRIMARY KEY,org_id TEXT,matter_id TEXT,
                  actor_user_id TEXT,action TEXT,target_type TEXT,target_id TEXT,details_json TEXT,created_at TEXT)''')
        c.commit()
        return c
    cipher = Fernet(Fernet.generate_key())
    monkeypatch.setattr(store, 'connect', connect)
    monkeypatch.setattr(store, 'encode', lambda x: cipher.encrypt(json.dumps(x, ensure_ascii=False).encode()))
    monkeypatch.setattr(store, 'decode', lambda x: json.loads(cipher.decrypt(x)))
    with connect() as c:
        c.executemany('INSERT INTO org_members VALUES (?,?,?)', [('o1','u1','owner'),('o1','reader','member'),('o2','u2','owner')])
    return path


def contract_payload():
    return {'name': 'secret-contract.txt', 'blocks': [{'id':'p1','text':'甲方应在验收后支付全部价款。','anchor':None}],
            'redacted_blocks':[{'id':'p1','text':'甲方应在验收后支付全部价款。'}], 'mapping':{}, 'warnings':[], 'format':'txt'}


def new_contract():
    return store.create_contract({'id':'m1','org_id':'o1'}, 'u1', contract_payload())


def test_redaction_keeps_material_terms():
    blocks = [{'id':'p1','text':'联系人：张三\n电话：13812345678；邮箱：test@example.com\n价款100万元，违约金30%，期限2026年9月25日。'}]
    result, mapping = docs.redact_blocks(blocks, ['张三'])
    assert '13812345678' not in result[0]['text']
    assert 'test@example.com' not in result[0]['text']
    assert '张三' not in result[0]['text']
    assert '100万元' in result[0]['text'] and '30%' in result[0]['text'] and '2026年9月25日' in result[0]['text']
    assert docs.restore(result[0]['text'], mapping) == blocks[0]['text']


def test_consistent_replacement_and_overlapping_terms():
    blocks = [{'id':'p1','text':'上海测试有限公司与上海测试有限公司；上海测试'}]
    result, mapping = docs.redact_blocks(blocks, ['上海测试','上海测试有限公司'])
    assert result[0]['text'].split('与')[0] == result[0]['text'].split('与')[1].split('；')[0]
    assert docs.restore(result[0]['text'], mapping) == blocks[0]['text']


def test_ordinary_role_clause_not_masked():
    text = '甲方：验收后支付全部价款。'
    result, _ = docs.redact_blocks([{'id':'p1','text':text}])
    assert result[0]['text'] == text


def word_file():
    d = Document(); d.add_paragraph('第一条 付款')
    d.add_paragraph('验收后90日付款。')
    d.add_table(rows=1, cols=1).cell(0,0).text = '附件：金额100万元。'
    b = io.BytesIO(); d.save(b); return b.getvalue()


def test_word_parse_and_anchored_revisions():
    original = word_file(); parsed = docs.parse_contract(original, 'a.docx')
    assert [b['text'] for b in parsed['blocks']] == ['第一条 付款', '验收后90日付款。', '附件：金额100万元。']
    updated = docs.redline_docx(original, parsed['blocks'], {'p2':'验收后60日付款。'})
    with zipfile.ZipFile(io.BytesIO(updated)) as z:
        xml = z.read('word/document.xml').decode()
        assert '<w:del ' in xml and '<w:ins ' in xml
        from lxml import etree
        root = etree.fromstring(z.read('word/document.xml'))
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        assert ''.join(root.xpath('//w:del//w:delText/text()', namespaces=ns)) == '9'
        assert ''.join(root.xpath('//w:ins//w:t/text()', namespaces=ns)) == '6'
        assert root.xpath('//w:p/w:r/w:t[text()="验收后"]', namespaces=ns)
        assert root.xpath('//w:p/w:r/w:t[text()="0日付款。"]', namespaces=ns)
        assert '附件：金额100万元。' in xml
    with pytest.raises(ValueError, match='未处理的修订'):
        docs.parse_contract(updated, 'a.docx')


def test_word_anchor_mismatch_fails_closed():
    original = word_file(); parsed = docs.parse_contract(original, 'a.docx')
    parsed['blocks'][1]['text'] = 'not the original'
    with pytest.raises(ValueError, match='定位已变化'):
        docs.redline_docx(original, parsed['blocks'], {'p2':'改动'})


@pytest.mark.parametrize('data,name', [(b'','a.txt'), (b'abc','a.exe'), (b'\xff','a.txt'), (b'x'*(docs.MAX_UPLOAD+1),'a.txt')])
def test_reject_invalid_uploads(data,name):
    with pytest.raises(ValueError): docs.parse_contract(data,name)


def test_no_silent_truncation():
    with pytest.raises(ValueError, match='不会静默截断'):
        docs.parse_contract(('条款'*40000).encode(), 'a.txt')


@pytest.mark.parametrize('url', ['https://www.gov.cn.evil.test/x','https://evil.test/www.gov.cn','http://www.gov.cn/x',
    'https://www.gov.cn@127.0.0.1/x','https://www.gov.cn:8080/x','https://127.0.0.1/x','file:///etc/passwd','javascript:alert(1)'])
def test_external_url_allowlist(url): assert not laws.official_url(url)


def test_official_url_allowed(): assert laws.official_url('https://www.court.gov.cn/fabu/xiangqing/419382.html')


def test_private_dns_is_rejected(monkeypatch):
    monkeypatch.setattr(laws.socket, 'getaddrinfo', lambda *a,**k: [(2,1,6,'',('127.0.0.1',443))])
    with pytest.raises(laws.LawRetrievalError): laws._public_host('www.gov.cn')


def test_external_search_failure_never_returns_law_memory(monkeypatch):
    monkeypatch.setenv('LEGAL_SEARCH_PROVIDER','bing_rss')
    def failed(request): return httpx.Response(503)
    with httpx.Client(transport=httpx.MockTransport(failed)) as client:
        result = laws.retrieve_law('民法典 违约金',['违约金'],client=client)
    assert result['status'] == 'unavailable' and result['sources'] == []


def test_external_source_refetched_and_quotes_not_search_snippets(monkeypatch):
    monkeypatch.setenv('LEGAL_SEARCH_PROVIDER','bing_rss'); monkeypatch.setattr(laws,'_public_host',lambda _:None)
    seen=[]
    def response(request):
        seen.append(str(request.url))
        if request.url.host == 'www.bing.com':
            return httpx.Response(200,text='<rss><channel><item><title>公开法规</title><link>https://www.gov.cn/test</link><description>搜索摘要不是证据</description></item></channel></rss>')
        return httpx.Response(200,headers={'content-type':'text/html; charset=utf-8'},text='<p>第一条 违约金规定。'+('可核验的测试正文。'*20)+'</p>')
    with httpx.Client(transport=httpx.MockTransport(response)) as client:
        result=laws.retrieve_law('测试法规',['违约金'],client=client)
    assert result['status']=='retrieved'
    assert len(seen)==2 and '搜索摘要不是证据' not in result['sources'][0]['text']
    assert result['sources'][0]['version_status']=='needs_verification'


def test_external_redirect_to_private_blocked(monkeypatch):
    monkeypatch.setattr(laws,'_public_host',lambda _:None)
    with httpx.Client(transport=httpx.MockTransport(lambda r:httpx.Response(302,headers={'location':'https://127.0.0.1/secrets'}))) as client:
        with pytest.raises(laws.LawRetrievalError): laws._read(client,'GET','https://www.gov.cn/x',official=True)


def candidate(kind='legal'):
    return {'rule_id':'liability','block_id':'p1','original_quote':'不承担任何赔偿责任。','title':'该条款违法无效',
            'kind':kind,'severity':'high','reason':'根据伪造第9999条违法。','impact':'违法',
            'suggested_text':'承担责任。','citations':[{'source_id':'made-up','supporting_quote':'伪造原文'}],'policy_ids':[]}


def validate(item, sources=None,policies=None):
    return engine.validate_result({'findings':[item]},[engine.RULES[2]],
        [{'id':'p1','text':'不承担任何赔偿责任。'}],sources or [],policies or [])


def test_fabricated_law_is_never_published_as_a_conclusion():
    result=validate(candidate()); f=result['findings'][0]
    assert f['evidence_status']=='unverified' and not f['suggested_text']
    assert '9999' not in json.dumps(f,ensure_ascii=False)
    assert f['title'].startswith('待核实')
    assert result['coverage'][0]['status']=='not_reviewed'


def test_exact_legal_quote_is_required():
    item=candidate(); item['citations']=[{'source_id':'law1','supporting_quote':'这是一条需要逐字匹配的测试法条原文。'}]
    source={'id':'law1','text':'这是一条需要逐字匹配的测试法条原文。'}
    assert validate(item,[source])['findings'][0]['evidence_status']=='source_matched'
    source['text']='内容不同'
    assert validate(item,[source])['findings'][0]['evidence_status']=='unverified'


def test_bad_contract_anchor_is_rejected():
    item=candidate(); item['block_id']='p999'
    assert validate(item)['findings']==[]
    item=candidate(); item['original_quote']='不在原文'
    assert validate(item)['rejected_findings']==1


def test_missing_clause_not_given_fake_anchor():
    item=candidate('commercial'); item.update(block_id=None,original_quote='')
    assert validate(item)['findings'][0]['suggested_text']==''


def test_company_policy_cannot_be_invented():
    assert validate(candidate('company_policy'))['findings']==[]
    item=candidate('company_policy'); item['policy_ids']=['p1']
    assert len(validate(item,policies=[{'id':'p1'}])['findings'])==1


def test_encrypted_storage_and_tenant_rules(db):
    cid=new_contract(); assert store.contract(cid)['payload']['name']=='secret-contract.txt'
    p=store.save_policy('o1','u1',{'title':'内部秘密规则','text':'内部规范正文','contract_type':'全部'})
    with pytest.raises(HTTPException) as e: store.policies('o1','u2')
    assert e.value.status_code==403
    with pytest.raises(HTTPException): store.save_policy('o1','reader',{'title':'x'})
    assert store.policies('o2','u2')==[]
    assert b'secret-contract.txt' not in db.read_bytes()
    assert '内部规范正文'.encode() not in db.read_bytes()
    assert store.policies('o1','reader')[0]['id']==p['id']


def test_policy_compare_and_swap(db):
    p=store.save_policy('o1','u1',{'title':'v1','text':'rule1','contract_type':'全部'})
    assert store.save_policy('o1','u1',{'title':'v2','text':'rule2','contract_type':'全部'},p['id'],1)['version']==2
    with pytest.raises(HTTPException) as e: store.save_policy('o1','u1',{'title':'overwrite'},p['id'],1)
    assert e.value.status_code==409


def test_redaction_is_explicit_and_immutable(db):
    cid=new_contract(); c=store.contract(cid)
    assert c['status']=='redaction_pending'
    store.confirm_redaction(cid,'u1',1,c['payload']['redacted_blocks'],{})
    assert store.contract(cid)['status']=='ready'
    with pytest.raises(HTTPException): store.confirm_redaction(cid,'u1',1,[],{})


def test_job_idempotency_and_durable_claim(db):
    c=store.contract(new_contract())
    rid,created=store.create_review(c,'u1','fp',{'stage':'new','policies':[]})
    assert created and store.create_review(c,'u1','fp',{})==(rid,False)
    token=store.claim(rid); assert token and store.claim(rid) is None
    store.checkpoint(rid,token,{'stage':'checkpoint','batches':{'1':{'completed':True}}},'failed')
    assert store.review(rid)['payload']['batches']['1']['completed']
    assert store.claim(rid) is not None
    with pytest.raises(RuntimeError): store.checkpoint(rid,token,{})


def test_decision_version_and_conflicting_edits(db):
    c=store.contract(new_contract()); rid,_=store.create_review(c,'u1','fp',{})
    findings=[{'id':'f1','block_id':'p1','kind':'commercial','evidence_status':'not_applicable'},
              {'id':'f2','block_id':'p1','kind':'commercial','evidence_status':'not_applicable'}]
    token=store.claim(rid); store.checkpoint(rid,token,{'findings':findings},'completed')
    assert store.decide(rid,'u1','f1','accepted','改后完整段落',0)['version']==1
    with pytest.raises(HTTPException): store.decide(rid,'u1','f2','accepted','冲突修改',0)
    with pytest.raises(HTTPException): store.decide(rid,'u1','f1','rejected','',0)
    store.decide(rid,'u1','f1','pending','',1)
    assert store.decide(rid,'u1','f2','accepted','改后段落',0)['version']==1


def test_unverified_legal_change_cannot_be_accepted(db):
    c=store.contract(new_contract()); rid,_=store.create_review(c,'u1','fp',{})
    token=store.claim(rid); store.checkpoint(rid,token,{'findings':[{'id':'f1','block_id':'p1','kind':'legal','evidence_status':'unverified'}]},'completed')
    with pytest.raises(HTTPException): store.decide(rid,'u1','f1','accepted','修改',0)


def ready_review():
    cid = new_contract()
    c = store.contract(cid)
    store.confirm_redaction(cid, 'u1', 1, c['payload']['blocks'], {})
    return store.create_review(store.contract(cid), 'u1', 'checkpoint-test',
                              {'contract_revision': 1, 'profile': {'our_role': '采购方'}, 'policies': []})[0]


def test_worker_resume_keeps_completed_batches(db, monkeypatch):
    rid = ready_review()
    monkeypatch.setattr(engine, 'authorize_job', lambda *args: None)
    searched, evaluated = [], []
    def search(query, keywords):
        searched.append(query)
        return {'status': 'retrieved', 'sources': [], 'provider': 'test', 'warnings': []}
    monkeypatch.setattr(laws, 'retrieve_law', search)
    failed_once = [False]
    def model(system, payload):
        ids = [r['id'] for r in payload['rules']]
        evaluated.append(ids)
        if ids[0] == 'termination' and not failed_once[0]:
            failed_once[0] = True
            raise RuntimeError('injected timeout')
        return {'coverage': [{'rule_id': i, 'status': 'reviewed', 'note': 'fixture'} for i in ids], 'findings': []}
    monkeypatch.setattr(engine, 'model_json', model)
    engine.run_review(rid)
    assert store.review(rid)['status'] == 'failed'
    assert len(store.review(rid)['payload']['batches']) == 1
    engine.run_review(rid)
    assert store.review(rid)['status'] == 'completed'
    assert len(searched) == len(engine.RULES)
    assert [x[0] for x in evaluated] == ['capacity', 'termination', 'termination']


def test_revoked_access_stops_before_any_external_call(db, monkeypatch):
    rid = ready_review()
    def deny(*args): raise HTTPException(403, 'revoked')
    monkeypatch.setattr(engine, 'authorize_job', deny)
    monkeypatch.setattr(laws, 'retrieve_law', lambda *args: pytest.fail('must not send data'))
    monkeypatch.setattr(engine, 'model_json', lambda *args: pytest.fail('must not call model'))
    engine.run_review(rid)
    assert store.review(rid)['status'] == 'failed'
    assert not store.review(rid)['payload'].get('searches')


def test_empty_searches_result_in_partial_not_clearance(db, monkeypatch):
    rid = ready_review()
    monkeypatch.setattr(engine, 'authorize_job', lambda *args: None)
    monkeypatch.setattr(laws, 'retrieve_law', lambda *args: {'status':'unavailable','sources':[],'provider':'test'})
    monkeypatch.setattr(engine, 'model_json', lambda system,p: {'coverage':[{'rule_id':r['id'],'status':'reviewed'} for r in p['rules']], 'findings':[]})
    engine.run_review(rid)
    assert store.review(rid)['status'] == 'partial'


def test_search_redirect_cannot_leak_authorization():
    with httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(302, headers={'location':'https://evil.test/x'}))) as client:
        with pytest.raises(laws.LawRetrievalError):
            laws._read(client, 'POST', 'https://api.tavily.com/search', headers={'Authorization':'Bearer test'}, json={'query':'public law'})


def test_api_requires_permission_redaction_and_consent(db, monkeypatch):
    router_module = pytest.importorskip('api.routers.contract_review')
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from configs import settings
    actor = {'id':'u1'}
    application = FastAPI()
    application.include_router(router_module.router)
    application.dependency_overrides[router_module.get_current_user] = lambda: actor
    def member(mid, user, min_role='viewer', require_active=False):
        if user['id'] != 'u1' or mid != 'm1': raise HTTPException(403, 'no access')
        return {'matter':{'id':'m1','org_id':'o1'}, 'role':'lead'}
    monkeypatch.setattr(router_module.matters, 'require_matter_member', member)
    monkeypatch.setattr(settings, 'openai_configured', lambda: True)
    submitted = []
    monkeypatch.setattr(engine, 'submit', submitted.append)
    client = TestClient(application)
    profile = {'our_role':'采购方','contract_type':'采购合同','external_processing_confirmed':True}
    upload = client.post('/legal/contracts', data={'matter_id':'m1'}, files={'file':('sample.txt','联系人：张三；付款100万元。'.encode(),'text/plain')})
    assert upload.status_code == 201
    cid = upload.json()['id']
    assert 'original_b64' not in upload.text and '张三' not in upload.text
    assert client.post(f'/legal/contracts/{cid}/reviews', json=profile).status_code == 409
    assert client.post(f'/legal/contracts/{cid}/redaction', json={'revision':1,'confirmed':True}).status_code == 200
    assert client.post(f'/legal/contracts/{cid}/reviews', json={**profile,'external_processing_confirmed':False}).status_code == 409
    response = client.post(f'/legal/contracts/{cid}/reviews', json=profile)
    assert response.status_code == 202 and len(submitted) == 1
    assert client.post(f'/legal/contracts/{cid}/reviews', json=profile).json()['id'] == response.json()['id']
    assert len(submitted) == 1
    actor['id'] = 'u2'
    assert client.get(f'/legal/contracts/{cid}').status_code == 403
    assert client.get('/legal/reviews/'+response.json()['id']).status_code == 403

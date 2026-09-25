"""Regression checks for Word accept/reject semantics and original document fidelity."""
import io
import zipfile
from copy import deepcopy
import pytest
from docx import Document
from docx.shared import Pt
from lxml import etree
from legal.docx_redlines import NS, W, open_package, paragraph_text, redline_docx


def package(doc):
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def blocks(data):
    archive, root = open_package(data)
    try:
        return [{'id': f'p{i+1}', 'text': paragraph_text(p).strip(), 'anchor': i}
                for i, p in enumerate(root.xpath('//w:body//w:p', namespaces=NS)) if paragraph_text(p).strip()]
    finally:
        archive.close()


def part(data, name='word/document.xml'):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return etree.fromstring(z.read(name))


def view(data, mode):
    root = deepcopy(part(data))
    for tag in ('ins', 'del'):
        for node in list(root.xpath(f'//w:{tag}', namespaces=NS)):
            parent, at = node.getparent(), node.getparent().index(node)
            if (mode == 'accept' and tag == 'ins') or (mode == 'reject' and tag == 'del'):
                for child in list(node):
                    for text in child.xpath('.//w:delText', namespaces=NS):
                        text.tag = f'{{{W}}}t'
                    parent.insert(at, child)
                    at += 1
            parent.remove(node)
    return [paragraph_text(p) for p in root.xpath('//w:body//w:p', namespaces=NS)]


def fixture():
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = '合同审查测试 · 内部样例'
    doc.add_heading('采购合同（测试样例）', 0)
    p = doc.add_paragraph()
    p.add_run('第一条 付款：验收后 ')
    run = p.add_run('90')
    run.bold = True
    run.font.size = Pt(12)
    p.add_run(' 日内付款。')
    doc.add_paragraph('  第二条\t交货：\n按附件安排。  ')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '项目'
    table.cell(0, 1).text = '保修期 12 个月'
    doc.sections[0].footer.paragraphs[0].text = '未签署的测试文件'
    return package(doc)


def test_actual_revisions_reconstruct_original_and_proposed_and_keep_format():
    data = fixture()
    original = blocks(data)
    replacement = {'p2': '第一条 付款：验收后 30 日内付款。', 'p5': '保修期 24 个月'}
    revised = redline_docx(data, original, replacement)
    assert view(revised, 'reject') == view(data, 'reject')
    expected = view(data, 'accept')
    expected[1], expected[4] = replacement['p2'], replacement['p5']
    assert view(revised, 'accept') == expected
    root = part(revised)
    assert root.xpath('//w:ins/w:r/w:rPr/w:b', namespaces=NS)
    assert root.xpath('//w:del/w:r/w:rPr/w:b', namespaces=NS)
    assert not root.xpath('//w:del//w:t', namespaces=NS)
    assert root.xpath('//w:p/w:r/w:t[text()=" 日内付款。"]', namespaces=NS)
    ids = root.xpath('//w:ins/@w:id|//w:del/@w:id', namespaces=NS)
    assert len(ids) == len(set(ids))
    assert all(x.get(f'{{{W}}}author') and x.get(f'{{{W}}}date') for x in root.xpath('//w:ins|//w:del', namespaces=NS))
    settings = part(revised, 'word/settings.xml')
    assert settings.find(f'{{{W}}}trackRevisions') is not None
    assert settings.find(f'{{{W}}}revisionView').get(f'{{{W}}}insDel') == 'true'
    with zipfile.ZipFile(io.BytesIO(data)) as old, zipfile.ZipFile(io.BytesIO(revised)) as new:
        for name in old.namelist():
            if name not in ('word/document.xml', 'word/settings.xml', '[Content_Types].xml', 'word/_rels/document.xml.rels'):
                assert old.read(name) == new.read(name), name


@pytest.mark.parametrize('before,after', [
    ('甲方付款。', '甲方分期付款。'), ('甲方应当无条件付款。', '甲方应当付款。'),
    ('Payment within 90 days.', 'Payment within 30 days.'), ('A\tB\nC', 'A\tD\nC'),
    ('  原有文字  ', '新的文字'), ('abc', ''), ('支付100元及50元', '支付120元及40元'),
])
def test_insert_delete_replace_breaks_and_boundary_whitespace(before, after):
    doc = Document()
    doc.add_paragraph(before)
    data = package(doc)
    revised = redline_docx(data, blocks(data), {'p1': after})
    assert view(revised, 'reject') == [before]
    prefix, suffix = before[:len(before)-len(before.lstrip())], before[len(before.rstrip()):]
    assert view(revised, 'accept') == [prefix + after + suffix]


def test_no_change_is_byte_identical():
    data = fixture()
    assert redline_docx(data, blocks(data), {}) == data
    assert redline_docx(data, blocks(data), {'p2': blocks(data)[1]['text']}) == data


def rewrite(data, callback):
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, 'w') as dst:
        for name in src.namelist():
            payload = callback(name, src.read(name))
            if payload is not None:
                dst.writestr(name, payload)
    return out.getvalue()


def test_settings_are_added_when_absent():
    doc = Document()
    doc.add_paragraph('90 天')
    data = rewrite(package(doc), lambda n, v: None if n == 'word/settings.xml' else v)
    revised = redline_docx(data, blocks(data), {'p1': '30 天'})
    assert part(revised, 'word/settings.xml').find(f'{{{W}}}trackRevisions') is not None


@pytest.mark.parametrize('tag', ['bookmarkStart', 'hyperlink', 'fldSimple', 'commentRangeStart'])
def test_complex_changed_paragraph_fails_without_flattening(tag):
    doc = Document()
    p = doc.add_paragraph('原文')
    marker = etree.SubElement(p._p, f'{{{W}}}{tag}')
    marker.set(f'{{{W}}}id', '100')
    doc.add_paragraph('其他段落')
    data = package(doc)
    with pytest.raises(ValueError):
        redline_docx(data, blocks(data), {'p1': '修改'})
    revised = redline_docx(data, blocks(data), {'p2': '其他内容'})
    assert etree.tostring(part(data).xpath('//w:p', namespaces=NS)[0]) == etree.tostring(part(revised).xpath('//w:p', namespaces=NS)[0])


def test_unknown_and_stale_anchors_fail():
    data = fixture()
    with pytest.raises(ValueError):
        redline_docx(data, blocks(data), {'unknown': '修改'})
    original = blocks(data)
    original[1]['text'] = '不同版本'
    with pytest.raises(ValueError):
        redline_docx(data, original, {'p2': '修改'})


def test_existing_revision_is_never_silently_accepted():
    data = fixture()
    revised = redline_docx(data, blocks(data), {'p2': '修改'})
    with pytest.raises(ValueError, match='已有未处理的修订'):
        open_package(revised)


def test_intake_and_export_use_the_same_anchor_text():
    from legal import contract_documents as documents
    data = fixture()
    parsed = documents.parse_contract(data, 'test.docx')
    assert parsed['blocks'][2]['text'] == '第二条\t交货：\n按附件安排。'
    revised = documents.redline_docx(data, parsed['blocks'], {'p3': '第二条\t交货：\n五日内交货。'})
    assert view(revised, 'reject') == view(data, 'reject')
    assert view(revised, 'accept')[2] == '  第二条\t交货：\n五日内交货。  '
